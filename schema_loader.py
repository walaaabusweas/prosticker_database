from pathlib import Path
from typing import Any
import openpyxl

SCHEMA_SHEETS = {
    'FINAL_TYPES', 'PRODUCT_SUBTYPES', 'PROPERTIES_BY_TYPE', 'SUBTYPE_PROPERTIES',
    'PRINTING_METHODS', 'OCCASIONS', 'CUSTOMIZATION_METHODS', 'PRICING_METHODS'
}


def _rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(x).strip() if x is not None else '' for x in rows[0]]
    return [dict(zip(headers, r)) for r in rows[1:] if any(x is not None and str(x).strip() != '' for x in r)]


def load_schema(path: str | Path) -> dict[str, Any]:
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    missing = SCHEMA_SHEETS - set(wb.sheetnames)
    if missing:
        raise ValueError(f'Missing schema sheets: {sorted(missing)}')

    final_types = _rows(wb['FINAL_TYPES'])
    subtypes = _rows(wb['PRODUCT_SUBTYPES'])
    common = _rows(wb['PROPERTIES_BY_TYPE'])
    subtype_props = _rows(wb['SUBTYPE_PROPERTIES'])

    properties: list[dict[str, Any]] = []
    for r in common:
        properties.append({
            'final_type': r.get('Final Type'), 'subtype': None,
            'scope': r.get('Scope') or 'COMMON', 'property': r.get('Property'),
            'description_ar': r.get('الوصف بالعربي'), 'data_type': r.get('Data Type') or 'Text',
            'allowed_values': r.get('Example / Allowed Values'),
            'required': str(r.get('Required?') or '').strip().lower() == 'yes',
            'notes': r.get('Notes')
        })
    for r in subtype_props:
        properties.append({
            'final_type': r.get('Final Type'), 'subtype': r.get('Subtype'),
            'scope': r.get('Scope') or 'SUBTYPE_ONLY', 'property': r.get('Property'),
            'description_ar': r.get('الوصف بالعربي'), 'data_type': r.get('Data Type') or 'Text',
            'allowed_values': r.get('Example / Allowed Values'),
            'required': str(r.get('Required?') or '').strip().lower() == 'yes',
            'notes': r.get('Notes')
        })

    lookups = {}
    for sheet, key in [
        ('PRINTING_METHODS', 'printing_methods'),
        ('OCCASIONS', 'occasions'),
        ('CUSTOMIZATION_METHODS', 'customization_methods'),
        ('PRICING_METHODS', 'pricing_methods'),
    ]:
        lookups[key] = _rows(wb[sheet])

    return {
        'final_types': final_types,
        'subtypes': subtypes,
        'properties': properties,
        'lookups': lookups,
    }


def type_map(schema):
    return {r['Final Type']: r for r in schema['final_types'] if r.get('Final Type')}


def subtypes_for(schema, final_type):
    return [r for r in schema['subtypes'] if r.get('Final Type') == final_type]


def properties_for(schema, final_type, subtype=None):
    result = []
    seen = set()
    for r in schema['properties']:
        if r.get('final_type') != final_type:
            continue
        if r.get('scope') == 'COMMON' or r.get('subtype') is None:
            key = r['property']
            if key not in seen:
                result.append(r); seen.add(key)
    if subtype:
        for r in schema['properties']:
            if r.get('final_type') == final_type and r.get('subtype') == subtype:
                key = r['property']
                if key not in seen:
                    result.append(r); seen.add(key)
    return result
