"""
core.py  —  منطق قاعدة البيانات (بدون أي اعتماد على Streamlit)

- يقرأ السكيما من ملف Excel (المصدر الوحيد للحقيقة).
- يبني الحقول (dropdown / multi-select / number / ...) لكل Type + Subtype.
- يحفظ المنتجات في SQLite ويصدّرها.
"""
from __future__ import annotations

import ast
import json
import os
import re
import sqlite3
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
SCHEMA_FILE = Path(os.environ.get("SCHEMA_FILE", BASE_DIR / "Full_Printing_Packaging_Schema_REVISED.xlsx"))
DB_FILE = Path(os.environ.get("CATALOG_DB", BASE_DIR / "catalog.db"))

SEP = "; "                      # فاصل القيم المتعددة داخل أعمدة SQLite الأساسية
OTHER_TOKENS = ("custom", "other")  # اختيار هذه القيم يفتح خانة "تفاصيل"

# ----------------------------------------------------------------------------
# إعدادات قابلة للتعديل بسهولة
# ----------------------------------------------------------------------------

# الخصائص التي يصح منطقيًا اختيار أكثر من قيمة لها (المنتج يتوفر بعدة خيارات).
# أي خاصية لها قائمة خيارات ولا تظهر هنا = Dropdown باختيار واحد.
MULTI_PROPS = {
    "occasion", "customization_method", "finish", "design_file", "use_cases", "target_businesses",
    "usage", "printing_sides", "closure_type", "handle_type", "special_finish", "application_surface",
    "print_area", "attachment_method", "additional_options", "food_type", "barrier_type",
    "material_layers", "binding_type", "fold_type", "lamination", "cover_finish", "protection_finish",
    "inner_layout", "hole_position", "bag_color", "paper_color", "cup_color", "body_color", "pen_color",
    "color", "ink_color", "shape", "stamp_shape", "tag_shape", "coaster_shape", "brooch_shape",
    "seal_shape", "cup_size", "cup_capacity", "coffee_capacity", "fill_volume", "pizza_size",
    "napkin_size", "paper_size", "poster_size", "sheet_size", "envelope_size", "invitation_size",
    "brooch_size", "stamp_size", "coaster_size", "sticker_size", "document_size", "special_function",
    "special_closure", "display_type", "print_material", "tote_material", "fabric_type",
    "wristband_material", "pen_type", "calendar_type", "planner_period", "menu_format", "document_type",
    "document_purpose", "event_type", "contact_info", "sleeve_fit_product", "wrap_target",
    "closure_target", "kraft_grade", "vintage_style", "box_style", "bag_style", "folder_purpose",
    "item_category", "print_type", "indoor_outdoor", "custom_finishing", "finishing", "invoice_format",
}

# حقول نصية حرة دائمًا (حتى لو فيها مثال فيه "/")
FREE_TEXT_PROPS = {
    "product_name", "description", "material", "size", "notes", "production_time", "guest_name",
    "table_number", "seat_number", "isbn", "pricing_formula", "fit_dimensions", "display_dimensions",
    "handle_color", "stand_material", "ribbon_color", "available_colors", "item_material",
}
TEXTAREA_PROPS = {"description", "notes"}

PAPER_SIZES = ["A3", "A4", "A5", "A6", "custom"]
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September",
          "October", "November", "December"]
CURRENCIES = ["ILS", "USD", "JOD", "SAR"]

# قوائم مقترحة (يمكن تعديلها) — لأن السكيما لم تحدد لها قيمًا ثابتة
DEFAULT_USE_CASES = ["تغليف منتجات", "هدايا", "فعاليات", "تسويق وإعلان", "طباعة مستندات",
                     "ضيافة وتقديم طعام", "بيع بالتجزئة", "هوية وعلامة تجارية", "ديكور"]
DEFAULT_TARGET_BUSINESSES = ["مطاعم", "مقاهي", "متاجر", "شركات", "محلات حلويات ومخابز",
                             "محلات هدايا وورد", "مدارس وجامعات", "معارض وفعاليات", "متاجر إلكترونية"]

AR_LABELS = {
    "quantity_unit": {"piece": "قطعة", "sheet": "ورقة", "set": "طقم", "roll": "رول"},
    "bool": {"Yes": "نعم", "No": "لا"},
    "generic": {"other": "أخرى", "custom": "مخصص"},
}

# الأعمدة الأساسية في جدول products (= CLIENT_ENTRY_TEMPLATE بدون faq لأنه جدول مستقل)
CORE_COLUMNS = [
    "product_name", "product_type", "subtype", "description", "material", "size", "quantity",
    "quantity_unit", "minimum_quantity", "custom_design", "design_file", "printing_method",
    "customization_method", "printing_sides", "finish", "occasion", "pricing_type", "price",
    "price_currency", "production_time", "notes", "price_method", "pricing_formula",
    "formula_variables", "setup_fee", "use_cases", "target_businesses", "fixed_price", "quote_required",
]
NUM_COLUMNS = {"quantity", "minimum_quantity", "price", "setup_fee", "fixed_price", "quote_required"}
CORE_MULTI = {"design_file", "customization_method", "printing_sides", "finish", "occasion",
              "use_cases", "target_businesses"}
# الحقول التي تُرسم في أقسام ثابتة (وليست ضمن "الخصائص الخاصة")
CORE_FORM_PROPS = {
    "product_name", "description", "subtype", "material", "size", "quantity", "quantity_unit",
    "minimum_quantity", "custom_design", "design_file", "printing_method", "customization_method",
    "printing_sides", "finish", "occasion", "pricing_type", "price", "price_currency",
    "production_time", "notes", "faq", "use_cases", "target_businesses",
}
# حقول بمفاتيح ثابتة (لا تضيع قيمتها عند تغيير الـSubtype)
STABLE_FIELDS = {"product_name", "description", "material", "size", "quantity", "quantity_unit",
                 "minimum_quantity", "custom_design", "occasion", "production_time", "notes",
                 "use_cases", "target_businesses", "customization_method"}


# ----------------------------------------------------------------------------
# قراءة السكيما
# ----------------------------------------------------------------------------
@dataclass
class Schema:
    types: "OrderedDict[str, dict]"
    subtypes: dict
    type_props: dict
    sub_props: dict
    occasions: list
    printing: list
    custom_methods: list
    pricing: list
    catalog: list
    labels: dict


@dataclass
class FieldSpec:
    name: str
    label: str
    kind: str                    # text | textarea | number | bool | select | multi
    options: list = field(default_factory=list)
    required: bool = False
    level: str = "core"          # core | type | subtype
    stable: bool = False

    @property
    def other_tokens(self):
        return [o for o in self.options if o in OTHER_TOKENS]


def _rows(wb, name):
    ws = wb[name]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(next(it))]
    out = []
    for r in it:
        if r and any(c is not None for c in r):
            out.append({h: (str(v).strip() if isinstance(v, str) else v) for h, v in zip(header, r)})
    return out


def _meta(r):
    return {
        "label": r.get("الوصف بالعربي") or r.get("Property"),
        "dtype": r.get("Data Type") or "",
        "example": r.get("Example / Allowed Values"),
        "required": str(r.get("Required?") or "").lower() == "yes",
    }


def load_schema(path: Path = SCHEMA_FILE) -> Schema:
    wb = load_workbook(path, read_only=True, data_only=True)

    types = OrderedDict()
    for r in _rows(wb, "FINAL_TYPES"):
        types[r["Final Type"]] = {"ar": r.get("الاسم بالعربي") or "", "category": r.get("Category") or ""}

    subtypes = {}
    for r in _rows(wb, "PRODUCT_SUBTYPES"):
        lst = subtypes.setdefault(r["Final Type"], [])
        if r["Subtype"] not in lst:
            lst.append(r["Subtype"])

    type_props = {}
    for r in _rows(wb, "PROPERTIES_BY_TYPE"):
        type_props.setdefault(r["Final Type"], OrderedDict())[r["Property"]] = _meta(r)

    sub_props = {}
    for r in _rows(wb, "SUBTYPE_PROPERTIES"):
        sub_props.setdefault((r["Final Type"], r["Subtype"]), OrderedDict())[r["Property"]] = _meta(r)

    occasions = [(r["Value"], r.get("العربي") or r["Value"]) for r in _rows(wb, "OCCASIONS")]
    printing = [(r["Value"], r.get("العربي") or r["Value"]) for r in _rows(wb, "PRINTING_METHODS")]
    custom_methods = [(r["Value"], r.get("العربي") or r["Value"]) for r in _rows(wb, "CUSTOMIZATION_METHODS")]
    pricing = [(r["Value"], r.get("العربي") or r["Value"]) for r in _rows(wb, "PRICING_METHODS")]

    catalog = [r for r in _rows(wb, "CATALOG_ENTRIES") if r.get("Entry Kind") == "PRODUCT" and r.get("Final Type")]
    wb.close()

    labels = {
        "occasion": dict(occasions),
        "printing_method": dict(printing),
        "customization_method": dict(custom_methods),
        "price_method": dict(pricing),
        **{k: v for k, v in AR_LABELS.items()},
    }
    return Schema(types, subtypes, type_props, sub_props, occasions, printing, custom_methods, pricing,
                  catalog, labels)


# ----------------------------------------------------------------------------
# تحويل الأمثلة إلى خيارات
# ----------------------------------------------------------------------------
def parse_options(example) -> list:
    if example is None:
        return []
    s = str(example).strip()
    if not s or s.lower() in {"none", "optional"}:
        return []
    parts = [p.strip() for p in s.split(" / ") if p.strip()]
    # "20 / 25 / 30 / 32 cm"  ->  20 cm / 25 cm / 30 cm / 32 cm
    m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*([A-Za-z]+)", parts[-1]) if parts else None
    if m and all(re.fullmatch(r"\d+(?:\.\d+)?", p) for p in parts[:-1]):
        parts = [f"{p} {m.group(2)}" for p in parts[:-1]] + [parts[-1]]
    seen, out = set(), []
    for p in parts:
        if p not in seen:
            seen.add(p)
            out.append(p)
    return out


def _options_for(schema: Schema, name: str, meta: dict) -> list:
    if name == "occasion":
        return [v for v, _ in schema.occasions]
    if name == "customization_method":
        return [v for v, _ in schema.custom_methods]
    if name == "printing_method":
        allowed = [v for v, _ in schema.printing]
        opts = [o for o in parse_options(meta.get("example")) if o in allowed]
        return opts or [v for v in allowed if v != "uv_dtf"]
    if name == "quantity_unit":
        return parse_options(meta.get("example")) or list(AR_LABELS["quantity_unit"])
    if name == "price_currency":
        return CURRENCIES
    if name in ("open_size", "closed_size"):
        return PAPER_SIZES
    if name == "start_month":
        return MONTHS
    if name == "use_cases":
        return DEFAULT_USE_CASES + ["other"]
    if name == "target_businesses":
        return DEFAULT_TARGET_BUSINESSES + ["other"]
    return parse_options(meta.get("example"))


def make_spec(schema: Schema, name: str, meta: dict, level: str) -> FieldSpec:
    dtype = (meta.get("dtype") or "").lower()
    label = meta.get("label") or name
    required = bool(meta.get("required"))
    stable = name in STABLE_FIELDS

    if "boolean" in dtype:
        return FieldSpec(name, label, "bool", ["Yes", "No"], required, level, stable)
    if dtype == "number":
        return FieldSpec(name, label, "number", [], required, level, stable)
    if name in TEXTAREA_PROPS:
        return FieldSpec(name, label, "textarea", [], required, level, stable)

    opts = [] if name in FREE_TEXT_PROPS else _options_for(schema, name, meta)
    if len(opts) < 2:
        return FieldSpec(name, label, "text", [], required, level, stable)

    if name in MULTI_PROPS:
        if not any(o in OTHER_TOKENS for o in opts):
            opts = opts + ["other"]
        return FieldSpec(name, label, "multi", opts, required, level, stable)
    return FieldSpec(name, label, "select", opts, required, level, stable)


def build_specs(schema: Schema, ptype: str, subtype: str) -> dict:
    """يرجع {'core': {name: spec}, 'type': [spec], 'subtype': [spec]}"""
    props = OrderedDict()
    for p, m in schema.type_props.get(ptype, {}).items():
        props[p] = (m, "type")
    for p, m in schema.sub_props.get((ptype, subtype), {}).items():
        props[p] = (m, props[p][1] if p in props else "subtype")   # الأدق (Subtype) يغلب في الخيارات

    core = OrderedDict()
    always = {"product_name": {"label": "اسم المنتج", "dtype": "Text", "required": True},
              "description": {"label": "الوصف", "dtype": "Text"},
              "notes": {"label": "ملاحظات", "dtype": "Text"},
              "use_cases": {"label": "الاستخدامات", "dtype": "Text / List"},
              "target_businesses": {"label": "الأعمال/القطاعات المستهدفة", "dtype": "Text / List"}}
    order = ["product_name", "description", "material", "size", "custom_design", "design_file",
             "printing_method", "customization_method", "printing_sides", "finish", "occasion", "quantity",
             "quantity_unit", "minimum_quantity", "production_time", "use_cases", "target_businesses", "notes"]
    for name in order:
        if name in props:
            core[name] = make_spec(schema, name, props[name][0], "core")
        elif name in always:
            core[name] = make_spec(schema, name, always[name], "core")
    if "product_name" in core:
        core["product_name"].required = True

    lists = {"type": [], "subtype": []}
    for name, (meta, level) in props.items():
        if name in CORE_FORM_PROPS:
            continue
        lists[level].append(make_spec(schema, name, meta, level))
    return {"core": core, **lists}


def all_specs(specs: dict) -> list:
    return list(specs["core"].values()) + specs["type"] + specs["subtype"]


# ----------------------------------------------------------------------------
# تحويل بين حالة الواجهة والقيم المخزنة
# ----------------------------------------------------------------------------
def split_token(value: str):
    """'custom: 12x15' -> ('custom', '12x15')"""
    for t in OTHER_TOKENS:
        if value == t:
            return t, ""
        if value.startswith(t + ":"):
            return t, value[len(t) + 1:].strip()
    return None, ""


def value_to_state(spec: FieldSpec, value):
    """القيمة المخزنة -> (state, detail)"""
    if spec.kind in ("text", "textarea"):
        return ("" if value is None else str(value)), ""
    if spec.kind == "number":
        try:
            return (None if value in (None, "") else float(value)), ""
        except (TypeError, ValueError):
            return None, ""
    if spec.kind == "bool":
        return (value if value in ("Yes", "No") else ""), ""
    if spec.kind == "select":
        if value in (None, ""):
            return "", ""
        if value in spec.options:
            return value, ""
        tok, detail = split_token(str(value))
        if tok and tok in spec.options:
            return tok, detail
        return "", ""
    # multi
    if value in (None, ""):
        return [], ""
    items = value if isinstance(value, list) else [v.strip() for v in str(value).split(SEP) if v.strip()]
    out, detail = [], ""
    for it in items:
        if it in spec.options:
            out.append(it)
            continue
        tok, d = split_token(str(it))
        if tok and tok in spec.options:
            if tok not in out:
                out.append(tok)
            detail = detail or d
    return out, detail


def state_to_value(spec: FieldSpec, state, detail: str = ""):
    """حالة الواجهة -> القيمة التي تُخزَّن (None إن كانت فارغة)"""
    detail = (detail or "").strip().replace(SEP, ", ")
    if spec.kind in ("text", "textarea"):
        s = (state or "").strip()
        return s or None
    if spec.kind == "number":
        if state is None:
            return None
        f = float(state)
        return int(f) if f.is_integer() else f
    if spec.kind == "bool":
        return state if state in ("Yes", "No") else None
    if spec.kind == "select":
        if not state:
            return None
        if state in OTHER_TOKENS and detail:
            return f"{state}: {detail}"
        return state
    vals = []
    for v in state or []:
        vals.append(f"{v}: {detail}" if (v in OTHER_TOKENS and detail) else v)
    return vals or None


def is_empty(v) -> bool:
    return v is None or v == "" or v == []


# ----------------------------------------------------------------------------
# المعادلات
# ----------------------------------------------------------------------------
def analyze_formula(formula: str):
    """يرجع (variables, error). نتحقق من الصياغة فقط."""
    f = (formula or "").strip()
    if not f:
        return [], None
    f2 = f.replace("×", "*").replace("÷", "/")
    if not re.fullmatch(r"[\w\s+\-*/().,%]+", f2):
        return [], "المعادلة تحتوي رموزًا غير مسموحة"
    try:
        ast.parse(f2, mode="eval")
    except SyntaxError:
        return [], "صياغة المعادلة غير صحيحة"
    names = re.findall(r"\b[A-Za-z_]\w*\b(?!\s*\()", f2)
    seen, out = set(), []
    for n in names:
        if n not in seen:
            seen.add(n)
            out.append(n)
    return out, None


# ----------------------------------------------------------------------------
# SQLite
# ----------------------------------------------------------------------------
def get_conn(path: Path = DB_FILE) -> sqlite3.Connection:
    conn = sqlite3.connect(str(path))
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db(path: Path = DB_FILE):
    cols = []
    for c in CORE_COLUMNS:
        t = "REAL" if c in NUM_COLUMNS else "TEXT"
        nn = " NOT NULL" if c in ("product_name", "product_type") else ""
        cols.append(f"{c} {t}{nn}")
    ddl = (
        "CREATE TABLE IF NOT EXISTS products ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT, " + ", ".join(cols) +
        ", properties TEXT NOT NULL DEFAULT '{}', created_at TEXT, updated_at TEXT)"
    )
    with get_conn(path) as conn:
        conn.execute(ddl)
        conn.execute(
            "CREATE TABLE IF NOT EXISTS product_faq ("
            "id INTEGER PRIMARY KEY AUTOINCREMENT, "
            "product_id INTEGER NOT NULL REFERENCES products(id) ON DELETE CASCADE, "
            "question TEXT NOT NULL, answer TEXT)"
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_products_type ON products(product_type, subtype)")


def _to_db(col, v):
    if isinstance(v, list):
        return SEP.join(str(x) for x in v)
    return v


def split_record(values: dict):
    """values (name -> stored value) -> (core dict, properties dict)"""
    core, props = {}, {}
    for k, v in values.items():
        if is_empty(v):
            continue
        if k in CORE_COLUMNS:
            core[k] = _to_db(k, v)
        else:
            props[k] = v
    return core, props


def save_product(values: dict, faqs: list, product_id: int | None = None, path: Path = DB_FILE) -> int:
    core, props = split_record(values)
    now = datetime.now().isoformat(timespec="seconds")
    with get_conn(path) as conn:
        if product_id:
            sets = ", ".join(f"{c} = ?" for c in CORE_COLUMNS)
            conn.execute(
                f"UPDATE products SET {sets}, properties = ?, updated_at = ? WHERE id = ?",
                [core.get(c) for c in CORE_COLUMNS] + [json.dumps(props, ensure_ascii=False), now, product_id],
            )
            conn.execute("DELETE FROM product_faq WHERE product_id = ?", (product_id,))
            pid = product_id
        else:
            names = ", ".join(CORE_COLUMNS)
            q = ", ".join("?" for _ in CORE_COLUMNS)
            cur = conn.execute(
                f"INSERT INTO products ({names}, properties, created_at, updated_at) VALUES ({q}, ?, ?, ?)",
                [core.get(c) for c in CORE_COLUMNS] + [json.dumps(props, ensure_ascii=False), now, now],
            )
            pid = cur.lastrowid
        for qn, an in faqs:
            conn.execute("INSERT INTO product_faq (product_id, question, answer) VALUES (?, ?, ?)", (pid, qn, an))
    return pid


def delete_product(pid: int, path: Path = DB_FILE):
    with get_conn(path) as conn:
        conn.execute("DELETE FROM products WHERE id = ?", (pid,))


def get_product(pid: int, path: Path = DB_FILE):
    """يرجع (values dict, faqs list[(q,a)]) أو None"""
    with get_conn(path) as conn:
        row = conn.execute("SELECT * FROM products WHERE id = ?", (pid,)).fetchone()
        if not row:
            return None
        faqs = [(r["question"], r["answer"] or "") for r in
                conn.execute("SELECT question, answer FROM product_faq WHERE product_id = ? ORDER BY id", (pid,))]
    d = dict(row)
    props = json.loads(d.pop("properties") or "{}")
    d.pop("created_at", None)
    d.pop("updated_at", None)
    d.pop("id", None)
    values = {k: v for k, v in d.items() if v is not None}
    values.update(props)
    return values, faqs


def list_products(search: str = "", ptype: str = "", path: Path = DB_FILE) -> list:
    sql = ("SELECT id, product_name, product_type, subtype, printing_method, price_method, price, "
           "quantity, quantity_unit, updated_at FROM products WHERE 1=1")
    args = []
    if ptype:
        sql += " AND product_type = ?"
        args.append(ptype)
    if search:
        sql += " AND (product_name LIKE ? OR description LIKE ?)"
        args += [f"%{search}%", f"%{search}%"]
    sql += " ORDER BY id DESC"
    with get_conn(path) as conn:
        return [dict(r) for r in conn.execute(sql, args)]


def counts_by_type(path: Path = DB_FILE) -> dict:
    with get_conn(path) as conn:
        return {r["product_type"]: r["n"] for r in
                conn.execute("SELECT product_type, COUNT(*) n FROM products GROUP BY product_type")}


# ----------------------------------------------------------------------------
# التصدير
# ----------------------------------------------------------------------------
def export_records(path: Path = DB_FILE) -> list:
    """قائمة منتجات كاملة (properties مفتوحة + faq كقائمة) — مناسبة للشات بوت."""
    out = []
    with get_conn(path) as conn:
        for row in conn.execute("SELECT * FROM products ORDER BY id"):
            d = dict(row)
            props = json.loads(d.pop("properties") or "{}")
            faqs = [{"question": r["question"], "answer": r["answer"]} for r in
                    conn.execute("SELECT question, answer FROM product_faq WHERE product_id = ? ORDER BY id", (d["id"],))]
            rec = {k: v for k, v in d.items() if v is not None}
            for k in CORE_MULTI:
                if k in rec:
                    rec[k] = [x for x in rec[k].split(SEP) if x]
            rec["properties"] = props
            rec["faq"] = faqs
            out.append(rec)
    return out


def export_dataframe(path: Path = DB_FILE):
    import pandas as pd

    rows = []
    for rec in export_records(path):
        flat = {k: v for k, v in rec.items() if k not in ("properties", "faq")}
        for k in CORE_MULTI:
            if k in flat:
                flat[k] = SEP.join(flat[k])
        for k, v in rec["properties"].items():
            flat[k] = SEP.join(v) if isinstance(v, list) else v
        flat["faq"] = "\n".join(f"س: {f['question']} | ج: {f['answer']}" for f in rec["faq"])
        rows.append(flat)
    df = pd.DataFrame(rows)
    if df.empty:
        return pd.DataFrame(columns=["id"] + CORE_COLUMNS)
    front = [c for c in ["id"] + CORE_COLUMNS + ["faq"] if c in df.columns]
    rest = [c for c in df.columns if c not in front]
    return df[front + sorted(rest)]
